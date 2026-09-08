# -*- coding: utf-8 -*-
"""
抓取 Instagram 主页全部帖子数据并输出到 Excel（test.xlsx）。

数据源：OpenCLI 浏览器桥接（Chrome 已登录 Instagram 的会话）。
原理：在 instagram.com 页面内用 fetch 调用同源内部 API：
      GET /api/v1/feed/user/{username}/username/?count=12&max_id=... 分页拉取全部帖子，
      评论内容（可选）用 GET /api/v1/media/{media_id}/comments/?count=N 抓取。

输出列（19 列）：
  Report date | Post ID | Account ID | Account username | Account name |
  Description | Duration (sec) | Publish time | Permalink | Post type |
  Data comment | Date | Views | Likes | Shares | Comments | Saves | Reach | Follows

说明：
  - Saves / Reach / Follows 属于 Instagram Insights 私有指标，公开 API 无法获取，
    默认填 "N/A"。如需补充需 Meta Graph API access token 或登录为账户所有者抓取
    Insights 页面，可在此脚本基础上扩展。
  - Shares 取 IG 的转发数（media_repost_count）。

用法：
  python fetch_instagram_posts.py --users anycubicofficial,anycubic_deutschland
  python fetch_instagram_posts.py --users anycubic_deutschland --comments 1 --raw posts_raw.json
  python fetch_instagram_posts.py --users anycubicofficial --limit 24 --output quick_test.xlsx

前置条件：
  - Chrome + OpenCLI 扩展已连接（opencli doctor 通过）
  - Chrome 中已登录 Instagram（任意账号即可，公开帖子无需是目标账号所有者）
"""

import argparse
import json
import subprocess
import sys
import time
from datetime import datetime, date
from pathlib import Path

NODE = r'C:\Users\zhengjingyi\.workbuddy\binaries\node\versions\22.22.2\node.exe'
MAIN = (r'C:\Users\zhengjingyi\.workbuddy\binaries\node\versions\22.22.2'
        r'\node_modules\@jackwener\opencli\dist\src\main.js')
IG_APP_ID = '936619743392459'

# 输出列（与用户要求一致）
COLUMNS = [
    'Report date', 'Post ID', 'Account ID', 'Account username', 'Account name',
    'Description', 'Duration (sec)', 'Publish time', 'Permalink', 'Post type',
    'Data comment', 'Date', 'Views', 'Likes', 'Shares', 'Comments',
    'Saves', 'Reach', 'Follows',
]

# 需要 Insights 权限、公开 API 拿不到的指标
INSIGHTS_ONLY = {'Saves', 'Reach', 'Follows'}


def oc(*args, timeout=180):
    """Run opencli via node directly (bypasses cmd.exe quoting issues)."""
    r = subprocess.run([NODE, MAIN] + list(args), capture_output=True, text=True,
                       encoding='utf-8', errors='replace', timeout=timeout)
    return (r.stdout or r.stderr).strip()


def open_tab(session, url):
    """Open a URL in the browser session, return the tab id."""
    out = oc('browser', session, 'open', url)
    try:
        return json.loads(out).get('page')
    except (json.JSONDecodeError, TypeError):
        return None


def ev(session, tab, js, timeout=90):
    """Evaluate JS in the page, return parsed JSON value (or raw string)."""
    out = oc('browser', session, 'eval', js, '--tab', tab, timeout=timeout)
    try:
        return json.loads(out)
    except (json.JSONDecodeError, TypeError):
        return out


# ---------------------------------------------------------------------------
# 登录检查
# ---------------------------------------------------------------------------
def is_logged_in(session, tab):
    """Return True if the IG page session is authenticated."""
    js = ("JSON.stringify({has_session: (document.cookie.indexOf('sessionid') > -1),"
          " has_ds_user: (document.cookie.indexOf('ds_user_id') > -1),"
          " body: (document.body.innerText || '').slice(0, 120)})")
    info = ev(session, tab, js)
    if not isinstance(info, dict):
        return False
    if info.get('has_session') or info.get('has_ds_user'):
        return True
    # 未登录时 IG 首页会显示登录表单
    body = (info.get('body') or '')
    return 'Log in' not in body and '登录' not in body and len(body) > 40


def ensure_logged_in(session, tab):
    """Try opencli instagram login first; fall back to manual login guidance."""
    print('== checking Instagram login ==')
    try:
        out = oc('instagram', 'login', '--timeout', '15', '-f', 'json', timeout=60)
        try:
            data = json.loads(out)
            if data.get('status') == 'already_logged_in':
                print('  logged in as @%s (%s)'
                      % (data.get('username', ''), data.get('full_name', '')))
                return True
        except (json.JSONDecodeError, TypeError):
            pass
    except Exception as exc:  # noqa: BLE001
        print('  opencli instagram login failed: %s' % exc)
    if is_logged_in(session, tab):
        print('  browser session logged in')
        return True
    print('ERROR: Instagram not logged in. Please log in to instagram.com in the '
          'browser, then re-run.', file=sys.stderr)
    return False


# ---------------------------------------------------------------------------
# IG 内部 API（页面内 fetch，自动携带 cookie / csrf）
# ---------------------------------------------------------------------------
def ig_get_json(session, tab, path, retries=2, delay=2.0):
    """Fetch a same-origin IG API path from inside the page. Returns parsed JSON."""
    js = ("(async function(){"
          "var m = document.cookie.match(/csrftoken=([^;]+)/);"
          "var csrf = m ? m[1] : '';"
          "var r = await fetch('%s', {headers: {'x-csrftoken': csrf,"
          " 'x-ig-app-id': '%s', 'x-requested-with': 'XMLHttpRequest'}});"
          "var t = await r.text();"
          "try { return JSON.stringify({ok:true, status:r.status, data: JSON.parse(t)}); }"
          "catch(e){ return JSON.stringify({ok:false, status:r.status, data:t.slice(0,400)}); }"
          "})()") % (path, IG_APP_ID)
    for attempt in range(retries + 1):
        res = ev(session, tab, js)
        if isinstance(res, dict) and res.get('ok'):
            return res['data']
        if isinstance(res, dict) and res.get('status') == 200:
            return {}
        print('    ! fetch failed (attempt %d/%d): %s'
              % (attempt + 1, retries + 1, res if isinstance(res, str) else res))
        time.sleep(delay)
    return None


def get_account_info(session, tab, username):
    """Return {pk, username, full_name} for a username via the feed API."""
    data = ig_get_json(session, tab,
                       '/api/v1/feed/user/%s/username/?count=1' % username)
    if not data or not data.get('items'):
        return None
    user = data['items'][0].get('user') or {}
    return {'pk': user.get('pk') or user.get('id') or '',
            'username': user.get('username', username),
            'full_name': user.get('full_name', '')}


def fetch_all_posts(session, tab, username, limit=0, page_delay=0.5):
    """Paginate over the profile feed API and return all media items.

    limit=0 means all posts. Returns (items, account).
    """
    account = get_account_info(session, tab, username)
    if not account or not account['pk']:
        print('ERROR: cannot resolve user id for @%s (private account?)' % username,
              file=sys.stderr)
        return [], account

    items = []
    seen = set()
    max_id = None
    page = 0
    while True:
        path = '/api/v1/feed/user/%s/username/?count=12' % username
        if max_id:
            path += '&max_id=%s' % max_id
        data = ig_get_json(session, tab, path)
        if data is None:
            print('  ! page %d failed, stopping.' % page)
            break
        batch = data.get('items') or []
        new = 0
        for it in batch:
            pid = str(it.get('pk') or it.get('id') or '')
            if pid and pid not in seen:
                seen.add(pid)
                items.append(it)
                new += 1
        page += 1
        print('  page %d: +%d posts (total %d)'
              % (page, new, len(items)), flush=True)
        if not data.get('more_available'):
            break
        max_id = data.get('next_max_id') or data.get('profile_grid_items_cursor')
        if not max_id:
            break
        if limit and len(items) >= limit:
            items = items[:limit]
            break
        time.sleep(page_delay)
    return items, account


# ---------------------------------------------------------------------------
# 字段映射
# ---------------------------------------------------------------------------
def fmt_datetime(ts):
    try:
        return datetime.fromtimestamp(int(ts)).strftime('%Y-%m-%d %H:%M:%S')
    except (TypeError, ValueError, OSError):
        return ''


def map_post_type(item):
    product = item.get('product_type') or ''
    media_type = item.get('media_type')
    if product == 'clips':
        return 'REEL'
    if product == 'igtv':
        return 'IGTV'
    if product == 'feed_video' or media_type == 2:
        return 'VIDEO'
    if media_type == 8:
        return 'CAROUSEL'
    if media_type == 1:
        return 'IMAGE'
    return 'UNKNOWN'


def media_views(item):
    for key in ('play_count', 'content_views_count', 'view_count', 'ig_play_count'):
        v = item.get(key)
        if v not in (None, '', 0):
            return v
    return ''


def build_rows(items, username, account, report_date, comment_n=0,
               session=None, tab=None):
    """Map raw IG media items to the 19 output columns."""
    rows = []
    account_id = (account or {}).get('pk', '')
    account_username = (account or {}).get('username', username)
    account_name = (account or {}).get('full_name', '')
    for idx, it in enumerate(items, 1):
        post_id = str(it.get('pk') or '')
        publish = fmt_datetime(it.get('taken_at'))
        code = it.get('code') or ''
        ptype = map_post_type(it)
        if ptype == 'REEL':
            permalink = 'https://www.instagram.com/reel/%s/' % code if code else ''
        else:
            permalink = 'https://www.instagram.com/p/%s/' % code if code else ''
        caption = (it.get('caption') or {}).get('text') or ''
        description = caption.replace('\r', '')

        # 可选：抓前 N 条评论内容
        data_comment = ''
        if comment_n > 0 and post_id:
            cdata = ig_get_json(session, tab,
                                '/api/v1/media/%s/comments/?count=%d'
                                % (post_id, comment_n))
            if cdata and cdata.get('comments'):
                data_comment = '；'.join(
                    '@%s: %s' % (c.get('user', {}).get('username', ''),
                                 (c.get('text') or '').replace('\r', ''))
                    for c in cdata['comments'][:comment_n])

        row = {
            'Report date': report_date,
            'Post ID': post_id,
            'Account ID': account_id,
            'Account username': account_username,
            'Account name': account_name,
            'Description': description,
            'Duration (sec)': it.get('video_duration') or '',
            'Publish time': publish,
            'Permalink': permalink,
            'Post type': ptype,
            'Data comment': data_comment,
            'Date': publish[:10] if publish else '',
            'Views': media_views(it),
            'Likes': it.get('like_count') or 0,
            'Shares': it.get('media_repost_count') or 0,
            'Comments': it.get('comment_count') or 0,
        }
        for col in INSIGHTS_ONLY:
            row[col] = 'N/A'
        rows.append(row)
        if idx % 100 == 0:
            print('  mapped %d/%d' % (idx, len(items)), flush=True)
    return rows


# ---------------------------------------------------------------------------
# Excel 输出
# ---------------------------------------------------------------------------
def write_excel(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'IG Posts'

    header_fill = PatternFill('solid', fgColor='C13584')  # IG 品牌紫红
    header_font = Font(color='FFFFFF', bold=True)
    for col, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for r, row in enumerate(rows, start=2):
        for c, name in enumerate(COLUMNS, start=1):
            value = row.get(name, '')
            cell = ws.cell(row=r, column=c, value=value)
            if name in ('Description', 'Data comment'):
                cell.alignment = Alignment(wrap_text=True, vertical='top')
            elif isinstance(value, (int, float)):
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal='right')

    widths = {
        'Report date': 12, 'Post ID': 22, 'Account ID': 22,
        'Account username': 20, 'Account name': 20, 'Description': 55,
        'Duration (sec)': 13, 'Publish time': 20, 'Permalink': 45,
        'Post type': 11, 'Data comment': 45, 'Date': 12, 'Views': 12,
        'Likes': 10, 'Shares': 10, 'Comments': 10, 'Saves': 10,
        'Reach': 10, 'Follows': 10,
    }
    for c, name in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(name, 12)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(COLUMNS)),
                                      max(2, len(rows) + 1))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--users', default='anycubicofficial,anycubic_deutschland',
                    help='Comma-separated IG usernames')
    ap.add_argument('--output', default=str(Path(__file__).resolve().parent / 'test.xlsx'),
                    help='Output Excel file path')
    ap.add_argument('--raw', default='',
                    help='Optional JSON file to dump raw media items')
    ap.add_argument('--limit', type=int, default=0,
                    help='Max posts per account (0 = all)')
    ap.add_argument('--comments', type=int, default=0,
                    help='Fetch first N comments per post for Data comment column '
                         '(0 = skip, much faster)')
    ap.add_argument('--session', default='ar',
                    help='OpenCLI browser session name')
    ap.add_argument('--page-delay', type=float, default=0.5,
                    help='Delay seconds between API pages')
    ap.add_argument('--report-date', default=date.today().isoformat(),
                    help='Report date value (YYYY-MM-DD)')
    args = ap.parse_args()

    users = [u.strip() for u in args.users.split(',') if u.strip()]
    if not users:
        ap.error('--users must contain at least one username')

    session = args.session
    tab = open_tab(session, 'https://www.instagram.com/')
    if not tab:
        print('ERROR: cannot open browser tab. Is Chrome + OpenCLI extension running?',
              file=sys.stderr)
        sys.exit(1)
    print('tab: %s' % tab)
    time.sleep(6)
    if not ensure_logged_in(session, tab):
        sys.exit(1)

    all_rows = []
    raw_store = {}
    for username in users:
        print('\n== @%s ==' % username)
        items, account = fetch_all_posts(session, tab, username,
                                         limit=args.limit,
                                         page_delay=args.page_delay)
        print('  total posts: %d' % len(items))
        if not items:
            print('  !! no posts fetched for @%s' % username)
            continue
        raw_store[username] = items
        rows = build_rows(items, username, account, args.report_date,
                          comment_n=args.comments, session=session, tab=tab)
        all_rows.extend(rows)
        print('  rows built: %d' % len(rows))

    if args.raw:
        Path(args.raw).parent.mkdir(parents=True, exist_ok=True)
        with open(args.raw, 'w', encoding='utf-8') as f:
            json.dump(raw_store, f, ensure_ascii=False, indent=1)
        print('\nraw items saved -> %s' % args.raw)

    if not all_rows:
        print('ERROR: no data at all, nothing written.', file=sys.stderr)
        sys.exit(1)

    out = Path(args.output)
    write_excel(all_rows, out)
    print('\n[done] %d rows written -> %s' % (len(all_rows), out))


if __name__ == '__main__':
    main()
